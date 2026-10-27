#!/usr/bin/env python3
"""Build the FIVE GUYS review deck (5 slides, 16:9) with python-pptx.

    python make_presentation.py      ->  FiveGuys_Presentation.pptx

Slides carry only the title and the requested keywords, plus photos and icons.
Photos: Wikimedia Commons (assets/photos, credits in assets/CREDITS.md).
Icons:  Tabler Icons, MIT (assets/icons).
"""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent
PHOTOS = ROOT / "assets" / "photos"
ICONS = ROOT / "assets" / "icons"
OUTPUT = ROOT / "FiveGuys_Presentation.pptx"

# --- Brand palette: Five Guys red, white, black (+ neutral tints for cards)
RED = RGBColor(0xED, 0x1C, 0x24)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
TILE = RGBColor(0xF2, 0xF2, 0xF2)        # card fill on white slides
TILE_DARK = RGBColor(0x1E, 0x1E, 0x1E)   # card fill on black slides
RULE_DARK = RGBColor(0x3A, 0x3A, 0x3A)   # thin dividers on black
FONT = "Arial"

# --- Geometry (inches unless noted)
EMU_IN = 914400
SLIDE_W_EMU, SLIDE_H_EMU = 12192000, 6858000      # 13.333 x 7.5 in
SLIDE_W, SLIDE_H = SLIDE_W_EMU / EMU_IN, SLIDE_H_EMU / EMU_IN
CHECK_COLS, CHECK_ROWS = 60, 2                    # red/white checker band
CHECK_EMU = SLIDE_W_EMU // CHECK_COLS
BAND_Y = SLIDE_H - CHECK_ROWS * CHECK_EMU / EMU_IN
M = 0.6                                           # side margin
CONTENT_Y = 2.45                                  # content starts below the 2-line title

PHOTO_CREDITS = (
    "Photos (Wikimedia Commons): "
    "Slide 1 Andy Li, CC0; "
    "Slide 2 Tjeerd Wiersma, CC BY 2.0; "
    "Slide 3 Bobak Ha'Eri, CC BY 3.0; "
    "Slide 4 {bacon}; "
    "Slide 5 Simon Doggett, CC BY 2.0. "
    "Icons: Tabler Icons (MIT)."
)
BACON_CREDIT = "Sarah Stierch, CC0"


def E(inches):
    return Emu(int(round(inches * EMU_IN)))


# ---------------------------------------------------------------- text helpers
def run(text, size, color, bold=True, spc=None):
    return text, {"size": size, "color": color, "bold": bold, "spc": spc}


def fill_text(tf, paragraphs, align=PP_ALIGN.LEFT, line_spacing=None):
    """paragraphs: list of paragraphs, each a list of run() tuples."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for text, st in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name, f.size, f.bold = FONT, Pt(st["size"]), st["bold"]
            f.color.rgb = st["color"]
            if st["spc"] is not None:
                r._r.get_or_add_rPr().set("spc", str(int(st["spc"] * 100)))


def add_text(slide, x, y, w, h, paragraphs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=None, name=None):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tb.text_frame.vertical_anchor = anchor
    fill_text(tb.text_frame, paragraphs, align, line_spacing)
    if name:
        tb.name = name
    return tb


# --------------------------------------------------------------- shape helpers
def _flat(shape, color):
    """Solid fill, no outline, no theme style (so no inherited shadow/bevel)."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def add_box(slide, x, y, w, h, color, *, radius=0.0, kind=None, name=None):
    kind = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    _flat(s, color)
    if name:
        s.name = name
    return s


def tag(slide, x, y, w, h, runs, color=RED, name=None):
    """Rounded block with one centered line of text (used for keyword tags)."""
    s = add_box(slide, x, y, w, h, color, radius=0.14, name=name)
    fill_text(s.text_frame, [runs], align=PP_ALIGN.CENTER)
    s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return s


def add_fade(slide, x, y, w, h, hex_color, name="Photo fade"):
    """Rectangle that fades from solid `hex_color` (left) to transparent (right)."""
    s = add_box(slide, x, y, w, h, BLACK, name=name)
    solid = s._element.spPr.find(qn("a:solidFill"))
    grad = etree.fromstring(
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">'
        f'<a:gsLst><a:gs pos="0"><a:srgbClr val="{hex_color}"><a:alpha val="100000"/></a:srgbClr></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{hex_color}"><a:alpha val="0"/></a:srgbClr></a:gs></a:gsLst>'
        '<a:lin ang="0" scaled="0"/></a:gradFill>'
    )
    solid.addprevious(grad)
    solid.getparent().remove(solid)
    return s


def add_icon(slide, icon, color, x, y, size):
    pic = slide.shapes.add_picture(str(ICONS / f"{icon}-{color}.png"), E(x), E(y), E(size), E(size))
    pic.name = f"Icon {icon}"
    pic._element.nvPicPr.cNvPr.set("descr", f"{icon.replace('-', ' ')} icon")
    return pic


def icon_badge(slide, icon, x, y, d, circle, icon_color, scale=0.56):
    add_box(slide, x, y, d, d, circle, kind=MSO_SHAPE.OVAL, name=f"Badge {icon}")
    s = d * scale
    add_icon(slide, icon, icon_color, x + (d - s) / 2, y + (d - s) / 2, s)


def add_photo(slide, file, x, y, w, h, *, focus=(0.5, 0.5), alt=""):
    """Place a photo that covers the box (cropped around `focus`, never stretched)."""
    path = PHOTOS / file
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(str(path), E(x), E(y), E(w), E(h))
    box, img = w / h, iw / ih
    if img > box:
        keep = box / img
        left = min(max(focus[0] - keep / 2, 0), 1 - keep)
        pic.crop_left, pic.crop_right = left, 1 - keep - left
    else:
        keep = img / box
        top = min(max(focus[1] - keep / 2, 0), 1 - keep)
        pic.crop_top, pic.crop_bottom = top, 1 - keep - top
    pic.name = f"Photo {Path(file).stem}"
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def checker_band(slide, on=RED, off=WHITE):
    """Five Guys red/white checker pattern along the bottom edge (one group)."""
    grp = slide.shapes.add_group_shape()
    y0 = SLIDE_H_EMU - CHECK_ROWS * CHECK_EMU
    _flat(grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y0, SLIDE_W_EMU, CHECK_ROWS * CHECK_EMU), off)
    for r in range(CHECK_ROWS):
        for c in range(CHECK_COLS):
            if (r + c) % 2 == 0:
                _flat(grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, c * CHECK_EMU, y0 + r * CHECK_EMU,
                                           CHECK_EMU, CHECK_EMU), on)
    grp.name = "Checker band"


def new_slide(prs, color):
    s = prs.slides.add_slide(prs.slide_layouts[5])  # "Title Only"
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color
    return s


def header(slide, line1, line2, dark):
    """Two-line title (line 2 in red) at the same spot on every content slide."""
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = E(M), E(0.62), E(6.3), E(1.5)
    title.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    fill_text(title.text_frame,
              [[run(line1, 44, WHITE if dark else BLACK)], [run(line2, 44, RED)]],
              line_spacing=0.92)


# ---------------------------------------------------------------------- slides
def slide_title(prs):
    s = new_slide(prs, RED)
    px = 7.2
    add_photo(s, "store_brighton_night.jpg", px, 0, SLIDE_W - px, BAND_Y, focus=(0.46, 0.5),
              alt="Five Guys restaurant at night with its red neon FIVE GUYS sign")
    title = s.shapes.title
    title.left, title.top, title.width, title.height = E(M), E(1.68), E(6.45), E(2.45)
    title.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    fill_text(title.text_frame,
              [[run("FIVE GUYS", 84, WHITE, spc=-2)], [run("REVIEW", 84, BLACK, spc=-2)]],
              line_spacing=0.88)
    add_text(s, M, 4.42, 6.4, 1.1,
             [[run("Successful Marketing", 30, WHITE)], [run("& Signature Menu", 30, WHITE)]],
             name="Subtitle")
    checker_band(s)


def slide_marketing_1(prs):
    s = new_slide(prs, WHITE)
    header(s, "$0 Advertising", "& Open Kitchen", dark=False)
    px = 8.15
    add_photo(s, "open_kitchen_covent_garden.jpg", px, 0, SLIDE_W - px, BAND_Y, focus=(0.38, 0.5),
              alt="Five Guys staff in red shirts cooking burgers in an open kitchen")

    # $0 on TV ads: crossed-out TV icon + big $0
    top, height = CONTENT_Y, 4.2
    add_box(s, M, top, 2.7, height, RED, radius=0.15, name="Stat $0")
    add_icon(s, "device-tv-off", "white", M + 0.32, top + 0.35, 0.78)
    add_text(s, M + 0.25, top + 1.9, 2.4, 1.95, [[run("$0", 120, WHITE)]], anchor=MSO_ANCHOR.BOTTOM)

    items = [("messages", "Word-of-Mouth"), ("snowflake-off", "No Freezers"),
             ("leaf", "100% Fresh"), ("heart-handshake", "Customer Trust")]
    lx, lw, ih = 3.6, 4.05, 0.87
    gap = (height - len(items) * ih) / (len(items) - 1)
    for i, (icon, label) in enumerate(items):
        y = top + i * (ih + gap)
        hero = i == len(items) - 1          # trust = the result
        add_box(s, lx, y, lw, ih, RED if hero else TILE, radius=0.12, name=f"Keyword {label}")
        icon_badge(s, icon, lx + 0.15, y + 0.13, 0.61, WHITE if hero else RED,
                   "red" if hero else "white")
        add_text(s, lx + 0.95, y, lw - 1.0, ih, [[run(label, 26, WHITE if hero else BLACK)]],
                 anchor=MSO_ANCHOR.MIDDLE)
    checker_band(s)


def slide_marketing_2(prs):
    s = new_slide(prs, WHITE)
    header(s, "Mass Customization", "& Brand Experience", dark=False)
    px, ph = 7.3, 3.45
    add_photo(s, "free_peanuts.jpg", px, 0, SLIDE_W - px, ph, focus=(0.5, 0.55),
              alt="A box of free in-shell peanuts at a Five Guys restaurant")

    cy, ch = 3.8, 2.9
    cards = [  # x, w, icon, big, label, hero
        (M, 3.0, "burger", "15", "Free Toppings", False),
        (M + 3.3, 5.0, "arrows-shuffle", "250,000+", "Combinations", False),
        (M + 8.6, SLIDE_W - M - (M + 8.6), "gift", "FREE", "Peanuts", True),
    ]
    for x, w, icon, big, label, hero in cards:
        add_box(s, x, cy, w, ch, RED if hero else TILE, radius=0.15, name=f"Stat {label}")
        add_icon(s, icon, "white" if hero else "red", x + 0.3, cy + 0.35, 0.55)
        add_text(s, x + 0.27, cy + 0.98, w - 0.4, 1.15, [[run(big, 72, WHITE if hero else RED)]])
        add_text(s, x + 0.3, cy + 2.2, w - 0.4, 0.45, [[run(label, 24, WHITE if hero else BLACK)]])
    # 15 toppings -> 250,000+ combinations
    d = 0.58
    arrow = add_box(s, M + 3.15 - d / 2, cy + ch / 2 - d / 2, d, d, BLACK, kind=MSO_SHAPE.OVAL,
                    name="Arrow badge")
    fill_text(arrow.text_frame, [[run("→", 20, WHITE)]], align=PP_ALIGN.CENTER)
    arrow.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    checker_band(s)


def slide_burger(prs):
    s = new_slide(prs, BLACK)
    header(s, "Bacon", "Cheeseburger", dark=True)
    px = 6.5
    add_photo(s, "bacon_cheeseburger.jpg", px, 0, SLIDE_W - px, BAND_Y, focus=(0.5, 0.5),
              alt="Five Guys bacon cheeseburger")
    add_fade(s, px - 0.02, 0, 2.0, BAND_Y, "000000")

    # Culinary terms in red capitals
    rows = [
        [run("High-heat ", 26, WHITE), run("SEAR", 42, RED)],
        [run("CARAMELIZED", 42, RED), run(" Onions", 26, WHITE)],
        [run("SAUTÉED", 42, RED), run(" Mushrooms", 26, WHITE)],
    ]
    row_h, step = 0.85, 1.3
    top = CONTENT_Y + (BAND_Y - 0.3 - CONTENT_Y - (2 * step + row_h)) / 2
    for i, phrase in enumerate(rows):
        y = top + i * step
        if i:
            add_box(s, M, y - (step - row_h) / 2, 5.4, 0.014, RULE_DARK, name="Divider")
        add_text(s, M, y, 5.8, row_h, [phrase], anchor=MSO_ANCHOR.MIDDLE, name="Culinary term")
    checker_band(s)


def slide_fries(prs):
    s = new_slide(prs, BLACK)
    header(s, "Cajun Fries", "& Conclusion", dark=True)
    px = 6.9
    add_photo(s, "cajun_fries.jpg", px, 0, SLIDE_W - px, BAND_Y, focus=(0.4, 0.5),
              alt="Five Guys Cajun fries in a white cup")

    w = 5.3
    # 100% Peanut Oil
    y = CONTENT_Y
    add_box(s, M, y, w, 1.0, TILE_DARK, radius=0.14, name="Keyword peanut oil")
    icon_badge(s, "droplet", M + 0.18, y + 0.17, 0.66, RED, "white")
    add_text(s, M + 1.08, y, w - 1.2, 1.0, [[run("100% Peanut Oil", 28, WHITE)]],
             anchor=MSO_ANCHOR.MIDDLE)
    # BLANCH -> Double DEEP-FRY
    tag(s, M, y + 1.35, w, 0.95, [run("BLANCH", 36, WHITE)], name="Step blanch")
    add_text(s, M, y + 2.32, w, 0.6, [[run("↓", 30, WHITE)]], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    tag(s, M, y + 2.95, w, 0.95, [run("Double ", 24, WHITE, bold=False), run("DEEP-FRY", 36, WHITE)],
        name="Step deep-fry")
    checker_band(s)


# ---------------------------------------------------------------------- theme
def apply_brand_theme(prs):
    """Brand colors + Arial in the theme, 16:9-scaled master/layout placeholders."""
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    a = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", a)
    scheme.set("name", "Five Guys")
    colors = {"dk1": "000000", "lt1": "FFFFFF", "dk2": "1A1A1A", "lt2": "F2F2F2",
              "accent1": "ED1C24", "accent2": "000000", "accent3": "8C8C8C",
              "accent4": "B3141A", "accent5": "F47C80", "accent6": "595959",
              "hlink": "ED1C24", "folHlink": "B3141A"}
    for tag_name, val in colors.items():
        el = scheme.find(f"a:{tag_name}", a)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, qn("a:srgbClr")).set("val", val)
    fonts = root.find(".//a:fontScheme", a)
    fonts.set("name", "Five Guys")
    for which in ("majorFont", "minorFont"):
        fonts.find(f"a:{which}/a:latin", a).set("typeface", FONT)
    for effect in root.iterfind(".//a:effectStyle", a):  # flat look: no theme shadows/bevels
        for child in list(effect):
            effect.remove(child)
        etree.SubElement(effect, qn("a:effectLst"))
    theme.blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # The stock template is 4:3; stretch its placeholders to the 16:9 canvas.
    ratio = SLIDE_W_EMU / 9144000
    holders = list(prs.slide_master.placeholders)
    for layout in prs.slide_layouts:
        holders += list(layout.placeholders)
    for ph in holders:
        if ph._element.spPr.find(qn("a:xfrm")) is not None:
            ph.left, ph.width = Emu(int(ph.left * ratio)), Emu(int(ph.width * ratio))


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W_EMU), Emu(SLIDE_H_EMU)
    prs.part._element.find(qn("p:sldSz")).attrib.pop("type", None)  # template says 4:3
    apply_brand_theme(prs)
    props = prs.core_properties
    props.title = "FIVE GUYS Review - Successful Marketing & Signature Menu"
    props.subject = "English presentation"
    props.author = props.last_modified_by = ""
    props.comments = PHOTO_CREDITS.format(bacon=BACON_CREDIT)  # file info only, not on slides

    slide_title(prs)
    slide_marketing_1(prs)
    slide_marketing_2(prs)
    slide_burger(prs)
    slide_fries(prs)
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    build()
